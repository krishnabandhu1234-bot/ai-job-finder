"""Tests for raw text extraction (PDF/DOCX/TXT). PDF and DOCX fixtures are
generated on the fly with the same libraries the app uses to read them
(PyMuPDF, python-docx) - no binary fixture files checked into the repo,
and no external services required."""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from app.resume.parser import ResumeParseError, parse_resume_file

FIXTURES_DIR = Path(__file__).parent / "fixtures"
SAMPLE_TEXT = (FIXTURES_DIR / "sample_resume.txt").read_text(encoding="utf-8")


def _make_pdf(tmp_path: Path, text: str) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    # Insert as a simple text block; exact layout doesn't matter for these
    # tests, only that the text round-trips through extraction.
    page.insert_textbox(fitz.Rect(36, 36, 559, 800), text, fontsize=9)
    path = tmp_path / "resume.pdf"
    doc.save(path)
    doc.close()
    return path


def _make_docx(tmp_path: Path, text: str) -> Path:
    document = DocxDocument()
    for line in text.split("\n"):
        document.add_paragraph(line)
    path = tmp_path / "resume.docx"
    document.save(path)
    return path


def test_parse_txt_resume(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_text(SAMPLE_TEXT, encoding="utf-8")

    result = parse_resume_file(path)

    assert result.file_type == "txt"
    assert "Jordan Smith" in result.text
    assert "jordan.smith@example.com" in result.text
    assert result.warnings == []


def test_parse_pdf_resume(tmp_path):
    path = _make_pdf(tmp_path, SAMPLE_TEXT)

    result = parse_resume_file(path)

    assert result.file_type == "pdf"
    assert result.page_count == 1
    assert "Jordan Smith" in result.text
    assert "EXPERIENCE" in result.text


def test_parse_docx_resume(tmp_path):
    path = _make_docx(tmp_path, SAMPLE_TEXT)

    result = parse_resume_file(path)

    assert result.file_type == "docx"
    assert "Jordan Smith" in result.text
    assert "EDUCATION" in result.text


def test_parse_docx_extracts_table_content(tmp_path):
    document = DocxDocument()
    document.add_paragraph("Jane Doe")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Python"
    table.rows[0].cells[1].text = "AWS"
    path = tmp_path / "table_resume.docx"
    document.save(path)

    result = parse_resume_file(path)

    assert "Python" in result.text
    assert "AWS" in result.text


def test_empty_txt_file_produces_a_warning(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")

    result = parse_resume_file(path)

    assert result.text == ""
    assert result.warnings


def test_corrupt_pdf_raises_resume_parse_error(tmp_path):
    path = tmp_path / "corrupt.pdf"
    path.write_bytes(b"this is not a real pdf file")

    with pytest.raises(ResumeParseError):
        parse_resume_file(path)


def test_corrupt_docx_raises_resume_parse_error(tmp_path):
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"this is not a real docx file")

    with pytest.raises(ResumeParseError):
        parse_resume_file(path)


def test_unsupported_extension_raises_resume_parse_error(tmp_path):
    path = tmp_path / "resume.rtf"
    path.write_text("some text", encoding="utf-8")

    with pytest.raises(ResumeParseError):
        parse_resume_file(path)


def test_image_only_pdf_warns_about_no_extractable_text(tmp_path):
    doc = fitz.open()
    page = doc.new_page()
    # A page with no text inserted at all simulates a scanned/image-only PDF.
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 10, 10))
    pix.set_rect(pix.irect, (255, 255, 255))
    page.insert_image(fitz.Rect(0, 0, 100, 100), pixmap=pix)
    path = tmp_path / "scanned.pdf"
    doc.save(path)
    doc.close()

    result = parse_resume_file(path)

    assert result.text == ""
    assert any("scanned" in w.lower() or "no extractable" in w.lower() for w in result.warnings)
