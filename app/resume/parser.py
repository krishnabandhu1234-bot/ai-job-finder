"""Raw text extraction from resume files (section 2: PDF/DOCX/TXT).

This module's only job is bytes-on-disk -> plain text, plus surfacing
warnings for things that will hurt extraction quality (scanned/image-only
PDFs, password-protected files, empty files) without crashing the app.
Structured field extraction happens in extractor.py.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document as DocxDocument

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ResumeParseError(Exception):
    """Raised when a resume file cannot be read at all (corrupt, wrong
    format, password-protected with no accessible password, etc). The
    caller (the Resumes UI page) is expected to catch this, store the
    error message on the Resume row, and let the user try a different
    file - never crash the app over one bad upload."""


@dataclass
class ParsedDocument:
    text: str
    file_type: str
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)


def parse_resume_file(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".txt":
        return _parse_txt(path)
    raise ResumeParseError(
        f"Unsupported file type {suffix!r}. Supported formats: "
        f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}."
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise ResumeParseError(f"Could not open PDF (it may be corrupt or invalid): {exc}") from exc

    try:
        if doc.is_encrypted and not doc.authenticate(""):
            raise ResumeParseError(
                "This PDF is password-protected. Please upload an unprotected copy."
            )

        text_parts: list[str] = []
        warnings: list[str] = []
        for page in doc:
            try:
                text_parts.append(page.get_text())
            except Exception as exc:  # a single bad page shouldn't lose the rest
                warnings.append(f"Could not extract text from page {page.number + 1}: {exc}")

        text = "\n".join(text_parts).strip()
        if not text:
            warnings.append(
                "No extractable text was found. This PDF may be a scanned image "
                "rather than real text - OCR is not supported yet."
            )
        return ParsedDocument(text=text, file_type="pdf", page_count=doc.page_count, warnings=warnings)
    finally:
        doc.close()


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:
        raise ResumeParseError(f"Could not open DOCX (it may be corrupt or invalid): {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    # Some resumes lay out contact info / skills in tables rather than
    # paragraphs - don't silently drop that content.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    parts.append(cell_text)

    text = "\n".join(parts).strip()
    warnings = [] if text else ["No extractable text was found in this DOCX file."]
    return ParsedDocument(text=text, file_type="docx", warnings=warnings)


def _parse_txt(path: Path) -> ParsedDocument:
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_bytes().decode("utf-8", errors="replace")
    except OSError as exc:
        raise ResumeParseError(f"Could not read text file: {exc}") from exc

    text = text.strip()
    warnings = [] if text else ["This file is empty."]
    return ParsedDocument(text=text, file_type="txt", warnings=warnings)
